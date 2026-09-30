"""Build-only openpyxl template. It never includes customer or activity data."""
from pathlib import Path
from io import BytesIO
from base64 import b64encode, b64decode
from zipfile import ZipFile
import json, sys, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from openpyxl.utils import get_column_letter as L
from PIL import Image as PILImage

OUT=Path(sys.argv[1]) if len(sys.argv)>1 else Path('public')
OUT.mkdir(parents=True,exist_ok=True)
NAVY='123D5B'; BLUE='155F96'; TEAL='168B8A'; PALE='EFF6FA'; INK='182C3B'; MUTED='667585'; IMPORT='28744B'; ORANGE='B86A12'; PURPLE='7661A4'
wb=Workbook(); wb.remove(wb.active)
# Reuse the approved logo from the application, without recreating it.
html=(OUT/'pc.html').read_text(encoding='utf-8')
match=re.search(r'<img\b[^>]*class="logo"[^>]*src="data:image/[^;]+;base64,([^"]+)"',html)
if not match: raise RuntimeError('Approved GON logo not found in pc.html')
source=PILImage.open(BytesIO(b64decode(match.group(1)))).convert('RGBA')
source.thumbnail((460,112))
canvas=PILImage.new('RGB',(480,132),'white');canvas.paste(source,((480-source.width)//2,(132-source.height)//2),source)
logo=BytesIO();canvas.save(logo,format='PNG');logo_bytes=logo.getvalue()


def block(ws,rg,text,fill=None,size=11,bold=False,color=INK):
    ws.merge_cells(rg)
    for row in ws[rg]:
        for c in row:
            if fill: c.fill=PatternFill('solid',fgColor=fill)
    c=ws[rg.split(':')[0]]; c.value=text
    c.font=Font(name='Calibri',size=size,bold=bold,color=color)
    c.alignment=Alignment(vertical='center',wrap_text=True,indent=1)


def setup(name,widths,end):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False;ws.sheet_view.zoomScale=85
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.sheet_properties.outlinePr.summaryRight=False
    ws.sheet_properties.tabColor=BLUE if name=='Dashboard' else TEAL
    ws.sheet_format.defaultRowHeight=21
    for i,w in enumerate(widths,1): ws.column_dimensions[L(i)].width=w
    for row in ws.iter_rows(min_row=1,max_row=55 if name=='Dashboard' else 16,max_col=end):
        for c in row:c.font=Font(name='Calibri',size=11,color=MUTED);c.alignment=Alignment(vertical='center')
    for r in range(2,6):ws.row_dimensions[r].height=22
    if name=='Dashboard':
        block(ws,'B2:E5','',fill='FFFFFF');block(ws,'F2:P3','RENDICONTAZIONE ATTIVITÀ',NAVY,23,True,'FFFFFF');block(ws,'F4:P5','GON srl | Dashboard operativa',NAVY,12,False,'C4E4F3'); logo_pos='B2'
        block(ws,'B7:P7','Periodo e filtri',None,11,False,MUTED)
    else:
        block(ws,'A2:B5','',fill='FFFFFF');block(ws,f'C2:{L(end)}3',name.upper(),NAVY,22,True,'FFFFFF');block(ws,f'C4:{L(end)}5','GON srl | Rendicontazione attività',NAVY,11,False,'C4E4F3');logo_pos='A2'
        block(ws,f'A7:{L(end)}7','Periodo e filtri',None,11,False,MUTED)
    img=Image(BytesIO(logo_bytes));img.width=220;img.height=60;ws.add_image(img,logo_pos)
    ws.print_options.horizontalCentered=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.page_margins.left=.25;ws.page_margins.right=.25;ws.page_margins.top=.3;ws.page_margins.bottom=.35
    ws.oddFooter.left.text='GON srl | Uso aziendale';ws.oddFooter.right.text='Pagina &P di &N'
    ws.oddFooter.left.size=8;ws.oddFooter.right.size=8
    return ws


def table(ws,headers,nums=(),pct=(),dates=()):
    for i,h in enumerate(headers,1):
        c=ws.cell(8,i,h);c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(name='Calibri',size=10,bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[8].height=30
    for r in (9,10,11):
        ws.row_dimensions[r].height=32
        for i in range(1,len(headers)+1):
            c=ws.cell(r,i);c.fill=PatternFill('solid',fgColor='F1F6FA' if r==10 else 'FFFFFF');c.font=Font(name='Calibri',size=11,color=IMPORT if ws.title=='Attivita' else '000000');c.alignment=Alignment(wrap_text=True,vertical='center')
            if i in nums:c.number_format='0.00##';c.alignment=Alignment(horizontal='right',vertical='center')
            if i in pct:c.number_format='0.0%';c.alignment=Alignment(horizontal='right',vertical='center')
            if i in dates:c.number_format='dd/mm/yyyy'
            if r==11:
                c.fill=PatternFill('solid',fgColor='E0F0F0');c.font=Font(name='Calibri',size=11,bold=True,color='000000');c.border=Border(top=Side(style='thin',color=TEAL))
    ws.freeze_panes='D9' if ws.title=='Attivita' else 'B9';ws.auto_filter.ref=f'A8:{L(len(headers))}10';ws.print_title_rows='1:8';ws.print_area=f'A1:{L(len(headers))}12'

D=setup('Dashboard',[3]+[9]*15,16)
A=setup('Attivita',[13,24,30,15,42,58,12,13,19,16,25],11)
M=setup('Macroaree',[45,14,15,15,15,15,14,22],8)
C=setup('Clienti',[44,14,15,15,15,15,14,22,34],9)
table(A,['Data','Dipendente','Cliente / cantiere','Tipo ore','Macroarea','Descrizione','Ore','Origine','Stato','Ore timer originali','ID attività'],nums=(7,10),dates=(1,))
A.column_dimensions['K'].hidden=True;A.page_setup.paperSize=A.PAPERSIZE_A3
table(M,['Macroarea','Attività','Ore totali','Cantiere (h)','Viaggio (h)','Ufficio (h)','Quota ore','Ore da completare'],nums=(3,4,5,6,8),pct=(7,))
table(C,['Cliente / cantiere','Attività','Ore totali','Cantiere (h)','Viaggio (h)','Ufficio (h)','Quota ore','Ore da completare','Chiave cliente'],nums=(3,4,5,6,8),pct=(7,))
C.column_dimensions['I'].hidden=True
for rg,label in [('B9:D10','ORE REGISTRATE'),('F9:H10','ATTIVITÀ'),('J9:L10','CLIENTI'),('N9:P10','DA COMPLETARE')]:
    block(D,rg,label,'EEF5F8',10,True,MUTED)
for rg in ['B11:D13','F11:H13','J11:L13','N11:P13']:
    block(D,rg,0,'EEF5F8',30,True,INK)
D['B11'].number_format='0.00';D['N11'].font=Font(name='Calibri',size=30,bold=True,color=ORANGE)
block(D,'B16:J16','ORE PER MACROAREA',None,12,True,BLUE);block(D,'K16:P16','CANTIERE / VIAGGIO / UFFICIO',None,11,True,BLUE)
for r in range(17,37):D.row_dimensions[r].height=20
bar=BarChart();bar.type='bar';bar.style=10;bar.title=None;bar.legend=None;bar.x_axis.title='Ore';bar.y_axis.title=None;bar.gapWidth=60
bar.add_data(Reference(M,min_col=3,min_row=8,max_row=16),titles_from_data=True);bar.set_categories(Reference(M,min_col=1,min_row=9,max_row=16))
bar.series[0].graphicalProperties.solidFill=BLUE;bar.series[0].graphicalProperties.line.noFill=True
bar.dLbls=DataLabelList();bar.dLbls.showVal=True;bar.dLbls.position='outEnd'
bar.anchor=TwoCellAnchor(_from=AnchorMarker(col=1,row=17),to=AnchorMarker(col=10,row=36));D.add_chart(bar)
block(D,'B39:G39','TIPO ORE',NAVY,11,True,'FFFFFF');block(D,'H39:J39','ORE',NAVY,11,True,'FFFFFF');block(D,'K39:M39','QUOTA',NAVY,11,True,'FFFFFF');block(D,'N39:P39','ATTIVITÀ',NAVY,11,True,'FFFFFF')
for r,t in zip(range(40,43),['Cantiere','Viaggio','Ufficio']):
    block(D,f'B{r}:G{r}',t,PALE if r%2==0 else 'FFFFFF')
    for left,right in [('H','J'),('K','M'),('N','P')]:block(D,f'{left}{r}:{right}{r}',0,PALE if r%2==0 else 'FFFFFF',11,False,'000000')
    D[f'H{r}'].number_format='0.00';D[f'K{r}'].number_format='0.0%';D.row_dimensions[r].height=28
pie=DoughnutChart();pie.holeSize=68;pie.style=10;pie.title=None;pie.legend.position='b'
pie.add_data(Reference(D,min_col=8,min_row=40,max_row=42));pie.set_categories(Reference(D,min_col=2,min_row=40,max_row=42))
pie.series[0].data_points=[DataPoint(idx=i) for i in range(3)]
for pt,color in zip(pie.series[0].data_points,[BLUE,TEAL,PURPLE]):pt.graphicalProperties.solidFill=color;pt.graphicalProperties.line.noFill=True
pie.dLbls=DataLabelList();pie.dLbls.showPercent=True;pie.dLbls.position='bestFit'
pie.anchor=TwoCellAnchor(_from=AnchorMarker(col=10,row=17),to=AnchorMarker(col=16,row=36));D.add_chart(pie)
block(D,'B45:P46','Inclusione attività incomplete', 'FFF3DF',11,False,ORANGE)
block(D,'B48:P49','Le macroaree PROGRAMMATO / PROGRAMMATA restano distinte dalle macroaree ordinarie. I totali rappresentano le ore registrate nei filtri selezionati.',None,11,False,MUTED)
block(D,'B51:P52','Dati: archivio GON. Tutti i riepiloghi si riferiscono al foglio Attivita; nessun dato dimostrativo viene aggiunto alle esportazioni.',None,10,False,MUTED)
D.print_area='A1:P53';D.page_setup.fitToHeight=1
wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True
buf=BytesIO();wb.save(buf)
with ZipFile(BytesIO(buf.getvalue())) as z:
    parts={n:({'base64':b64encode(z.read(n)).decode()} if n.endswith('.png') else {'text':z.read(n).decode('utf-8')}) for n in z.namelist()}
(OUT/'assets').mkdir(parents=True,exist_ok=True)
(OUT/'assets/report-template.json').write_text(json.dumps({'version':'excel-report-1.0.0','parts':parts},ensure_ascii=False),encoding='utf-8')
print('EXCEL TEMPLATE: 4 sheets, 2 native charts, approved logo, no activity data')
